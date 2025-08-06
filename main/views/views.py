from django.shortcuts import render, get_object_or_404, redirect
from django.http import FileResponse, Http404, HttpResponse
from django.contrib.admin.views.decorators import staff_member_required
from datetime import datetime, timedelta
from docx import Document

from main.models.models import WeeklyJournal, Noticia, Categoría, WeeklyJournalForm

# --- Admin: List + Create ---
@staff_member_required
def admin_journals(request):
    if request.method == 'POST' and 'upload' in request.POST:
        form = WeeklyJournalForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('admin_journals')
    else:
        form = WeeklyJournalForm()

    journals = WeeklyJournal.objects.order_by('-uploaded_at')
    return render(request, 'news/admin_journals.html', {
        'form': form,
        'journals': journals,
    })

# --- Admin: Edit existing ---
@staff_member_required
def edit_journal(request, pk):
    journal = get_object_or_404(WeeklyJournal, pk=pk)
    if request.method == 'POST':
        form = WeeklyJournalForm(request.POST, request.FILES, instance=journal)
        if form.is_valid():
            form.save()
            return redirect('admin_journals')
    else:
        form = WeeklyJournalForm(instance=journal)

    return render(request, 'news/admin_edit_journal.html', {
        'form': form,
        'journal': journal,
    })

# --- Admin: Delete existing ---
@staff_member_required
def delete_journal(request, pk):
    journal = get_object_or_404(WeeklyJournal, pk=pk)
    if request.method == 'POST':
        journal.delete()
        return redirect('admin_journals')
    return render(request, 'news/admin_confirm_delete.html', {
        'journal': journal,
    })

# --- Public page: list + latest download ---
def public_journal_page(request):
    journals = WeeklyJournal.objects.order_by('-uploaded_at')
    return render(request, 'news/public_download.html', {
        'journals': journals,
    })

# --- Public download endpoint ---
def download_journal_file(request, pk):
    journal = get_object_or_404(WeeklyJournal, pk=pk)
    return FileResponse(
        journal.file.open('rb'),
        as_attachment=True,
        filename=journal.file.name.split('/')[-1]
    )

# --- Admin: Export weekly news document ---
@staff_member_required
def export_news_to_word(request):
    document = Document()
    document.add_heading('Noticias organizadas por categoría', level=1)

    today = datetime.today().date()
    days_since_sunday = (today.weekday() + 1) % 7
    last_sunday = today - timedelta(days=days_since_sunday)

    categories = Categoría.objects.all()
    for category in categories:
        news_items = Noticia.objects.filter(
            categoría=category,
            published_date__gte=last_sunday
        ).order_by('fecha_de_publicación')

        if news_items.exists():
            document.add_heading(category.name, level=1)
            for item in news_items:
                document.add_heading(item.título, level=2)
                document.add_paragraph(f"Fecha: {item.fecha_de_publicación.strftime('%d/%m/%Y')}")
                autores = ', '.join([a.nombre for a in item.autores.all()])
                document.add_paragraph(f"Autores: {autores}")
                document.add_paragraph(item.contenido)
                document.add_paragraph('---')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="noticias_por_categoria.docx"'
    document.save(response)
    return response

# --- Admin: Tools page ---
@staff_member_required
def admin_news_tools(request):
    return render(request, 'admin/news/export_news.html')
